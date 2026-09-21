#!/usr/bin/env python3
import sys
import base64
from docx import Document
from cryptography.hazmat.primitives.asymmetric import ed25519
from cryptography.hazmat.primitives import serialization

# Replace with the actual public key string from the TrustForge environment
PUBLIC_KEY_PEM = """-----BEGIN PUBLIC KEY-----
MCowBQYDK2VwAyEAV6KxJil7X0h/XydwAKHt6W7p94rF7U7n9/DIYqPjgQQ=
-----END PUBLIC KEY-----"""

def verify_docx(filepath):
    print(f"Verifying {filepath}...")
    try:
        doc = Document(filepath)
    except Exception as e:
        print(f"❌ Error reading DOCX: {e}")
        return False

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    
    # Extract text from tables to match the generator logic
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    table_texts.append(cell.text.strip())

    # Find the signature block
    sig_start = -1
    sig_end = -1
    for i, p in enumerate(paragraphs):
        if "--- BEGIN ED25519 SIGNATURE ---" in p:
            sig_start = i
        elif "--- END ED25519 SIGNATURE ---" in p:
            sig_end = i

    if sig_start == -1 or sig_end == -1 or sig_end <= sig_start:
        print("❌ Signature block not found in the document.")
        return False

    # Extract signature base64
    sig_b64 = "".join(paragraphs[sig_start + 1:sig_end]).replace("\n", "").strip()
    try:
        signature = base64.b64decode(sig_b64)
    except Exception:
        print("❌ Invalid base64 signature format.")
        return False

    # Reconstruct the signed payload
    # The payload was created from paragraphs BEFORE the signature block
    payload_paragraphs = paragraphs[:sig_start]
    
    text_payload = "\n".join(payload_paragraphs)
    if table_texts:
        text_payload += "\n" + "\n".join(table_texts)

    # Load public key
    try:
        public_key = serialization.load_pem_public_key(PUBLIC_KEY_PEM.encode('utf-8'))
    except Exception as e:
        print(f"❌ Failed to load public key: {e}")
        return False

    # Verify
    try:
        public_key.verify(signature, text_payload.encode('utf-8'))
        print("✅ SUCCESS: The document signature is VALID and the content has not been tampered with.")
        return True
    except Exception:
        print("❌ FAILURE: Signature verification failed. The document may have been tampered with.")
        return False

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python verify_audit.py <path_to_docx>")
        sys.exit(1)
    
    filepath = sys.argv[1]
    verify_docx(filepath)
