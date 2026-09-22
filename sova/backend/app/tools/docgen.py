"""FR-18/FR-19: DOCX deliverable generation + quality checks."""
import os
import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from ..core.config import settings


def generate_approval_note(task_id: str, subject: str, findings: str, sources: list,
                            calculations: str, recommendations: str, verification: dict,
                            compliance_table: list = None,
                            reviewer_name: str = "", decision: str = "",
                            decision_timestamp: str = "", comments: str = ""):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    for style_name in ("Heading 1", "Heading 2"):
        styles[style_name].font.name = "Aptos Display"
        styles[style_name].font.color.rgb = RGBColor(0x16, 0x3A, 0x5F)

    header = section.header.paragraphs[0]
    header.text = "TRUSTFORGE  |  INSPECTION REVIEW"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    footer = section.footer.paragraphs[0]
    footer.text = "Confidential - TrustForge local workbench | ⚠️ AI-Assisted Document: Human Verification Required"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    footer.runs[0].bold = True

    title = doc.add_heading("Inspection Approval Note", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("TrustForge secure AI workbench")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True
    metadata = doc.add_table(rows=2, cols=2)
    metadata.style = "Light Shading Accent 1"
    metadata.cell(0, 0).text = "Task ID"
    metadata.cell(0, 1).text = task_id
    metadata.cell(1, 0).text = "Generated"
    metadata.cell(1, 1).text = f"{datetime.datetime.utcnow().isoformat()} UTC"

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph("This note records the AI-assisted inspection review, its evidence, verification checks, and the required human decision.")

    doc.add_heading("1. Subject", level=1)
    doc.add_paragraph(subject or "Not specified")

    doc.add_heading("2. Extracted Findings", level=1)
    for paragraph in (findings or "No findings extracted.").split("\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())

    doc.add_heading("3. Relevant Evidence / Source References", level=1)
    if sources:
        for s in sources:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{s.get('filename')} — page {s.get('page')}").bold = True
            if s.get("chunk"):
                excerpt = " ".join(s["chunk"].split())[:240]
                doc.add_paragraph(f"Evidence excerpt: {excerpt}")
            if s.get("low_confidence"):
                r = p.add_run("  [LOW CONFIDENCE OCR]")
                r.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
    else:
        doc.add_paragraph("No source documents retrieved.")

    # Compliance comparison table
    if compliance_table:
        doc.add_heading("3b. Compliance Comparison", level=1)
        ct = doc.add_table(rows=len(compliance_table) + 1, cols=5)
        ct.style = "Light Grid Accent 1"
        headers = ["Parameter", "Measured", "Limit", "Status", "Source"]
        for j, h in enumerate(headers):
            ct.cell(0, j).text = h
            ct.cell(0, j).paragraphs[0].runs[0].bold = True if ct.cell(0, j).paragraphs[0].runs else False
        for i, row in enumerate(compliance_table, 1):
            ct.cell(i, 0).text = row.get("parameter", "")
            ct.cell(i, 1).text = row.get("measured", "")
            ct.cell(i, 2).text = row.get("limit", "")
            ct.cell(i, 3).text = row.get("status", "")
            ct.cell(i, 4).text = row.get("source_page", "")
            # Color the status cell
            if row.get("status") == "FAIL":
                for p in ct.cell(i, 3).paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)

    doc.add_heading("4. Calculations", level=1)
    doc.add_paragraph(calculations or "Not applicable.")

    doc.add_heading("5. Recommendations", level=1)
    doc.add_paragraph(recommendations or "None generated.")

    doc.add_heading("6. Verification Status", level=1)
    for k, v in (verification or {}).items():
        doc.add_paragraph(f"{k}: {v}", style="List Bullet")

    doc.add_heading("7. Review / Approval", level=1)
    if reviewer_name and decision:
        # Fill in approval info if available
        rows = [
            ("Reviewer Name", reviewer_name),
            ("Decision", decision),
            ("Date", decision_timestamp or datetime.datetime.utcnow().isoformat()),
            ("Comments", comments or "No comments."),
        ]
    else:
        rows = [("Reviewer Name", ""), ("Decision (Approve/Reject)", ""), ("Date", ""), ("Comments", "")]
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Light Grid Accent 1"
    for i, (k, v) in enumerate(rows):
        table.cell(i, 0).text = k
        table.cell(i, 1).text = v
        table.cell(i, 0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        table.cell(i, 1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    filename = f"Inspection_Approval_Note_{task_id[:8]}.docx"
    filepath = os.path.join(settings.ARTIFACT_DIR, filename)
    
    # 8. Sign the document
    if os.path.exists(settings.SIGNING_PRIVATE_KEY_PATH):
        from cryptography.hazmat.primitives.asymmetric import ed25519
        from cryptography.hazmat.primitives import serialization
        import base64
        
        # Extract text to sign
        text_payload = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_payload += "\n" + cell.text.strip()
        
        with open(settings.SIGNING_PRIVATE_KEY_PATH, "rb") as key_file:
            private_key = serialization.load_pem_private_key(
                key_file.read(),
                password=None
            )
        signature = private_key.sign(text_payload.encode('utf-8'))
        sig_b64 = base64.b64encode(signature).decode('utf-8')
        
        doc.add_paragraph("\n--- BEGIN ED25519 SIGNATURE ---")
        sig_p = doc.add_paragraph(sig_b64)
        sig_p.runs[0].font.name = "Courier New"
        sig_p.runs[0].font.size = Pt(8)
        sig_p.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        doc.add_paragraph("--- END ED25519 SIGNATURE ---")

    doc.save(filepath)
    return filepath, filename


def quality_check(filepath: str, required_sections=("Subject", "Findings", "Evidence", "Verification")):
    """FR-19: reopen and validate the generated file before calling it done."""
    try:
        doc = Document(filepath)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        missing = [s for s in required_sections if s.lower() not in full_text.lower()]
        opens_ok = True
    except Exception as e:
        return {"status": "FAIL", "reason": f"File does not open: {e}"}
    if missing:
        return {"status": "FAIL", "reason": f"Missing required sections: {missing}"}
    return {"status": "PASS", "reason": "All required sections present and file opens correctly."}
